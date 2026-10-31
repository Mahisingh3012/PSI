"""Builds PITCH.pptx (3 slides) from the run's media.  python scripts/make_pitch.py"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
MEDIA = ROOT / "runs/drishti/media"
INK, MUT = RGBColor(0x1F, 0x29, 0x37), RGBColor(0x6B, 0x72, 0x80)

p = Presentation()
p.slide_width, p.slide_height = Inches(13.33), Inches(7.5)


def slide(title, sub, bullets, img):
    s = p.slides.add_slide(p.slide_layouts[6])
    tf = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(1.1)).text_frame
    tf.text = title
    r = tf.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.color.rgb = Pt(30), True, INK
    q = tf.add_paragraph()
    q.text = sub
    q.runs[0].font.size, q.runs[0].font.color.rgb = Pt(15), MUT
    bf = s.shapes.add_textbox(Inches(0.6), Inches(1.7), Inches(6.2), Inches(5.4)).text_frame
    bf.word_wrap = True
    for i, b in enumerate(bullets):
        q = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
        q.text = b
        q.runs[0].font.size, q.runs[0].font.color.rgb = Pt(15), INK
        q.space_after = Pt(9)
    s.shapes.add_picture(str(MEDIA / img), Inches(7.0), Inches(1.8), width=Inches(5.9))


slide("The Case Study Engine", "From a messy 8-month project folder to a case study the client can sign", [
    "Agents with one job each: the LLM does the judgement (contract rules, KPI definitions, exposure, writing for three readers); code does the counting and checking.",
    "Every line cites a file and line, or a calculation. The verifier refuses wrong numbers, drifted citations, identifiers and unmarked estimates; a fixer agent gets the refused lines back.",
    "Two human stops (our delivery owner, then the client's authorised signatories). Approvals are bound to the pack's hash; 'no' and withdrawal are handled; every project is in one register.",
    "The same code runs another project's library: only a new metric spec (test 10)."],
    "diagram_onprem_architecture.png")
slide("What it found in the library", "All signed pilot KPIs recomputed exactly from raw data", [
    "Refused: the sales headline '74% to 97%, zero LTIs, Rs 2.3 Cr'. It mixes baselines, uses one zone's best week, and claims savings Finance rejected.",
    "True story: PPE 69% to 91% in the pilot (met), 94% in those areas by September; zone breaches -46% (missed), then -77% after the client's own guards.",
    "Credit shared: the client's safety week, toolbox talks, contractor discipline, guards, and phone escalation. Gate time is -24%, with the client's weighbridge and dispatch slots doing part.",
    "Exposure: plant GPS inside the snapshot file, signage in images, a public post linking client to vendor, and the plant's capacity in the annual report. All media was redrawn from data."],
    "chart_ppe_compliance_by_cohort.png")
slide("Getting to yes", "Safe and rewarding for the people who sign", [
    "Gatekeeper: on-prem, no egress, text-only alerts, face module removed with the hash verified by the client's IT, the outage handled openly.",
    "User: a normal day before and after. False alarms checked by their own team; the photo ended arguments; the gang leaders check helmets themselves.",
    "Buyer: one headline; money only where Finance measured it, with estimates marked ASSUMED. Champion: a five-line note to forward.",
    "Approval: to Corporate Comms + Legal by email, reply by 16 Nov (10 business days), per-asset options, and charts offered for their own council deck."],
    "chart_gate_turnaround_attribution.png")
p.save(ROOT / "PITCH.pptx")
print("PITCH.pptx written")
